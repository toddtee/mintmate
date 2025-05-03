# src/mintmate/excel_writer.py

import os
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from PIL import Image
from .puzzle_utils   import get_solution_san, get_to_move_color
from .screenshotter  import screenshot_puzzle

def create_workbook(*, puzzle_sets, df, output_dir, theme, style, file_naming, verbose=True):
    def log(msg):
        if verbose:
            print(msg)

    # Prepare naming and output directory
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M")
    override_fn = file_naming.get("override_filename", "")
    term        = file_naming.get("term", "")
    week        = file_naming.get("week", "")
    template    = file_naming.get("template", "mintmate_puzzles_{timestamp}.xlsx")
    cleanup     = file_naming.get("cleanup_screenshots", False)

    if override_fn:
        final_filename = override_fn
        dir_name       = os.path.splitext(override_fn)[0]
    else:
        no_ts = re.sub(r"[\s_-]*\{timestamp\}", "", template)
        final_filename = no_ts.format(term=term, week=week, timestamp=ts)
        dir_name       = os.path.splitext(final_filename)[0]

    batch_dir = os.path.join(output_dir, dir_name.replace(" ", ""))
    os.makedirs(batch_dir, exist_ok=True)
    log(f"📁 Output folder: {batch_dir}")

    # Create workbook and hidden AnswerKey
    wb        = Workbook()
    answer_ws = wb.active
    answer_ws.title = "AnswerKey"
    answer_ws.append(["Set","Puzzle ID","Answer"])
    answer_ws.sheet_state = "hidden"

    max_img_width = {}

    # ——— Generate one sheet per puzzle_set ———
    for ps in puzzle_sets:
        sheet_name = ps.get("sheet_name", "Puzzles")
        log(f"📑 Adding sheet: {sheet_name}")
        ws = wb.create_sheet(title=sheet_name)

        # Headers
        ws.append([
            "Puzzle ID","Type","Rating","To Move",
            "Position","Student Answer","Correct?","Lichess URL"
        ])

        # Folder for this set’s screenshots
        set_dir = os.path.join(batch_dir, sheet_name.replace(" ", ""))
        os.makedirs(set_dir, exist_ok=True)
        max_img_width[sheet_name] = 0

        # Populate puzzles
        for job in ps["puzzles"]:
            log(f"🎯 Generating {job['count']} '{job['type']}' puzzles [{job['min_rating']}–{job['max_rating']}] in '{sheet_name}'")
            subset = df[
                df["Themes"].str.contains(job["type"], case=False, na=False) &
                (df["Rating"] >= job["min_rating"]) &
                (df["Rating"] <= job["max_rating"])
            ]
            if subset.empty:
                log(f"⚠️ No puzzles for '{job['type']}' in {sheet_name}")
                continue

            for _ in range(job["count"]):
                p = subset.sample(1).iloc[0]
                pid, fen, moves = p["PuzzleId"], p["FEN"], p["Moves"]
                img_path = os.path.join(set_dir, f"{pid}.png")

                # Screenshot
                screenshot_puzzle(pid, img_path, theme, style, log)

                # AnswerKey
                sol = get_solution_san(fen, moves)
                answer_ws.append([sheet_name, pid, ", ".join(sol)])

                # To Move
                tm = get_to_move_color(fen, moves)
                url = f"https://lichess.org/training/{pid}?theme={theme}&piece={style}"

                # Write row
                row = ws.max_row + 1
                ws.append([pid, job["type"], p["Rating"], tm, "", "", "", url])
                log(f"✏️ Wrote row for {pid}")

                # Embed image and size row
                with Image.open(img_path) as img:
                    w, h = img.size
                max_img_width[sheet_name] = max(max_img_width[sheet_name], w)

                img_obj = XLImage(img_path)
                img_obj.width, img_obj.height = w, h
                ws.add_image(img_obj, f"E{row}")
                ws.row_dimensions[row].height = h * 0.85
                log(f"🖼️ Embedded screenshot for {pid}")

                # ✅/❌ formula
                ws[f"G{row}"] = (
                    f'=IF(F{row}="","",'
                    f'IF(F{row}=VLOOKUP(A{row},AnswerKey!$B:$C,2,FALSE),'
                    f'"✅ Correct","❌ Try again"))'
                )

    # ——— Remove default blank sheet ———
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    # ——— Auto-fit columns & center content ———
    log("🔧 Auto-fitting columns and centering")
    for sheet in wb.worksheets:
        if sheet.title == "AnswerKey":
            continue
        # measure text lengths
        col_lens = {}
        for cells in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in cells:
                length = len(str(cell.value or ""))
                col_lens[cell.column] = max(col_lens.get(cell.column, 0), length)

        for col, ln in col_lens.items():
            letter = get_column_letter(col)
            if letter == "E" and max_img_width.get(sheet.title):
                width = int(max_img_width[sheet.title] * 0.13) + 2
            elif letter == "G":
                width = max(len("✅ Correct"), len("❌ Try again")) + 2
            else:
                width = ln + 2
            sheet.column_dimensions[letter].width = width

        # center all cells
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # ——— Save workbook ———
    out_file = os.path.join(batch_dir, final_filename)
    wb.save(out_file)
    log(f"📄 Workbook saved to: {out_file}")

    # ——— Optional cleanup ———
    if cleanup:
        log("🗑️ Removing screenshots...")
        for root, _, files in os.walk(batch_dir):
            for f in files:
                if f.lower().endswith(".png"):
                    os.remove(os.path.join(root, f))
        log("✅ Screenshots cleaned up")
    else:
        log("📸 Screenshots retained")

    return out_file
