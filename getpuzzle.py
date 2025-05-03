import tomli
import pandas as pd
import os
import re
import chess
from datetime import datetime
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from PIL import Image
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import time

# === Helpers ===

def load_config(path="mintmate_config.toml"):
    with open(path, "rb") as f:
        return tomli.load(f)


def get_solution_san(original_fen, uci_moves):
    board = chess.Board(original_fen)
    ucs = uci_moves.split()
    # discard opponent setup move
    if ucs:
        first = chess.Move.from_uci(ucs[0])
        if first not in board.legal_moves:
            raise ValueError(f"Illegal setup move {ucs[0]} on FEN {original_fen}")
        board.push(first)
    san_list = []
    for u in ucs[1:]:
        move = chess.Move.from_uci(u)
        if move not in board.legal_moves:
            raise ValueError(f"Illegal puzzle move {u} in position {board.fen()}")
        san_list.append(board.san(move))
        board.push(move)
    return san_list


def get_to_move_color(original_fen, uci_moves):
    board = chess.Board(original_fen)
    ucs = uci_moves.split()
    if ucs:
        setup = chess.Move.from_uci(ucs[0])
        if setup in board.legal_moves:
            board.push(setup)
    return "White" if board.turn else "Black"


def screenshot_puzzle(puzzle_id, output_path, board_theme, piece_style, log):
    log(f"📸 Capturing screenshot for puzzle {puzzle_id}")
    url = f"https://lichess.org/training/{puzzle_id}?theme={board_theme}&piece={piece_style}"
    options = Options()
    options.add_argument("--headless")
    options.add_argument("--window-size=800,800")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    driver = webdriver.Chrome(options=options)
    driver.get(url)
    try:
        WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "cg-board"))
        )
        time.sleep(1)
        driver.execute_script("document.querySelectorAll('square.last-move').forEach(el => el.remove());")
        board = driver.find_element(By.CSS_SELECTOR, "cg-board")
        board.screenshot(output_path)
        log(f"✅ Screenshot saved to: {output_path}")
    except Exception as e:
        print(f"❌ Screenshot failed: {e}")
    finally:
        driver.quit()

# === Main ===
if __name__ == "__main__":
    cfg = load_config()
    # verbosity control
    verbose = cfg.get("verbose", True)
    def log(msg):
        if verbose:
            print(msg)

    # load config
    log("🔄 Loading configuration")
    csv_path = cfg["csv_path"]
    log(f"🔢 Loading puzzle data from {csv_path}")
    df = pd.read_csv(csv_path)
    log(f"✅ Loaded {len(df)} puzzles")

    # setup output dirs
    base_output = cfg["output_dir"]
    log(f"📂 Ensuring base output directory: {base_output}")
    os.makedirs(base_output, exist_ok=True)

    board_theme = cfg["board_theme"]
    piece_style = cfg["piece_style"]
    naming = cfg.get("file_naming", {})
    override_fn = naming.get("override_filename")
    term = naming.get("term", "")
    week = naming.get("week", "")
    template = naming.get("template", "mintmate_puzzles_{timestamp}.xlsx")
    cleanup = naming.get("cleanup_screenshots", False)

    # determine filenames & dirs
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M")
    if override_fn:
        final_filename = override_fn
        dir_name = os.path.splitext(override_fn)[0]
    else:
        no_ts = re.sub(r"[\s_-]*\{timestamp\}", "", template)
        final_filename = no_ts.format(term=term, week=week, timestamp=ts)
        dir_name = os.path.splitext(final_filename)[0]
    out_dir = os.path.join(base_output, dir_name.replace(" ", ""))
    log(f"📁 Creating output folder: {out_dir}")
    os.makedirs(out_dir, exist_ok=True)

    puzzle_sets = cfg.get("puzzle_sets") or [{"sheet_name": "Foundation", "puzzles": cfg.get("puzzles", [])}]
    wb = Workbook()

    # setup answer key
    log("🔐 Creating hidden AnswerKey sheet")
    answer_ws = wb.create_sheet(title="AnswerKey")
    answer_ws.append(["Set","Puzzle ID","Answer"])
    answer_ws.sheet_state = "hidden"

    max_img_width = 0

    # generate sheets
    for idx, ps in enumerate(puzzle_sets):
        sheet_name = ps.get("sheet_name", f"Set{idx+1}")
        if idx == 0:
            ws = wb.active
            ws.title = sheet_name
        else:
            log(f"📑 Adding sheet: {sheet_name}")
            ws = wb.create_sheet(title=sheet_name)
        ws.append(["Puzzle ID","Type","Rating","To Move","Position","Student Answer","Correct?","Lichess URL"])

        set_dir = os.path.join(out_dir, sheet_name.replace(" ", ""))
        log(f"📂 Creating folder for set '{sheet_name}': {set_dir}")
        os.makedirs(set_dir, exist_ok=True)

        for job in ps["puzzles"]:
            log(f"🎯 Generating {job['count']} '{job['type']}' puzzles [{job['min_rating']}–{job['max_rating']}] in '{sheet_name}'")
            subset = df[
                df["Themes"].str.contains(job["type"], case=False, na=False) &
                (df["Rating"] >= job["min_rating"]) &
                (df["Rating"] <= job["max_rating"])
            ]
            if subset.empty:
                log(f"⚠️ No puzzles found for type '{job['type']}' in {sheet_name}")
                continue

            for _ in range(job["count"]):
                p = subset.sample(1).iloc[0]
                pid, fen, moves = p["PuzzleId"], p["FEN"], p["Moves"]
                img_path = os.path.join(set_dir, f"{pid}.png")

                screenshot_puzzle(pid, img_path, board_theme, piece_style, log)
                sol = get_solution_san(fen, moves)
                answer_ws.append([sheet_name, pid, ", ".join(sol)])

                tm = get_to_move_color(fen, moves)
                url = f"https://lichess.org/training/{pid}?theme={board_theme}&piece={piece_style}"

                row = ws.max_row + 1
                ws.append([pid, job["type"], p["Rating"], tm, "", "", "", url])
                log(f"✏️ Writing row for puzzle {pid} in sheet '{sheet_name}'")

                with Image.open(img_path) as img:
                    w, h = img.size
                max_img_width = max(max_img_width, w)
                img_obj = XLImage(img_path)
                img_obj.width, img_obj.height = w, h
                ws.add_image(img_obj, f"E{row}")
                log(f"🖼️ Embedded screenshot for puzzle {pid}")

                ws.row_dimensions[row].height = h * 0.85
                ws.cell(row=row, column=7).value = (
                    f'=IF(F{row}="","",IF(F{row}=VLOOKUP(A{row},AnswerKey!$B:$C,2,FALSE),"✅ Correct","❌ Try again"))'
                )

    log("✍️ Auto-fitting columns and centering text...")
    for sheet in wb.worksheets:
        if sheet.title == "AnswerKey":
            continue
        col_lens = {}
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                length = len(str(cell.value or ""))
                col_lens[cell.column] = max(col_lens.get(cell.column, 0), length)
        if max_img_width:
            col_lens[5] = int(max_img_width * 0.13)
        corr = max(len("✅ Correct"), len("❌ Try again"))
        col_lens[7] = corr
        for col, length in col_lens.items():
            sheet.column_dimensions[get_column_letter(col)].width = length + 2
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                cell.alignment = Alignment(horizontal='center', vertical='center')

    out_file = os.path.join(out_dir, final_filename)
    wb.save(out_file)
    log(f"📄 Workbook exported to: {out_file}")

    if cleanup:
        log("🗑️ Cleaning up screenshots...")
        for root, dirs, files in os.walk(out_dir):
            for f in files:
                if f.lower().endswith('.png'):
                    os.remove(os.path.join(root, f))
        log("✅ Screenshots removed")
    else:
        log("📸 Screenshots retained in puzzle folder")
